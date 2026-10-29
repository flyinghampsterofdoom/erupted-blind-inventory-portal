"""Versioned, bounded XLSX protocol. Never evaluate formulas or fetch links.

Application runtime reuses the accepted Phase5 pinned workbook engine. Artifact
qualification separately imports/renders the actual server-generated originals.
"""
from io import BytesIO
from decimal import Decimal,InvalidOperation
from zipfile import ZipFile,BadZipFile,ZipInfo,ZIP_DEFLATED
from xml.etree import ElementTree as ET
import hashlib,hmac,re
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment,Protection
from openpyxl.worksheet.datavalidation import DataValidation
from .identity import digest

VERSION='erupted-quotation/1'
MAX_FILE=10*1024*1024
MAX_EXPANDED=64*1024*1024
MAX_LINES=5000
MAX_ENTRIES=128
START=12
HEADERS=['Line identity','Our SKU','Vendor SKU','Product','Variation','Units per pack','Requested packs','Vendor can fill packs','Vendor price per pack (USD)']
MANIFEST=['protocol','request_id','vendor_id','revision_id','revision_number','currency','snapshot_digest','signature']
PUBLIC={'protocol','request_id','vendor_id','revision_id','revision_number','currency','supplier','purchaser','ship_to','lines'}
LINE={'line_id','sku','vendor_sku','product','variation','pack_size','requested_packs'}


class InvalidWorkbook(ValueError):
 def __init__(self,errors):self.errors=errors if isinstance(errors,list) else [str(errors)];super().__init__('; '.join(self.errors[:12]))


def projection(request_id,revision_id,number,snapshot):
 p=dict(protocol=VERSION,request_id=request_id,vendor_id=snapshot['supplier_id'],revision_id=revision_id,revision_number=number,currency=snapshot['currency'],supplier=snapshot['supplier']['name'],purchaser=snapshot['header']['purchaser'],ship_to=snapshot['header']['ship_to'],lines=[dict(line_id=l['row_id'],sku=l['sku'],vendor_sku=l['vendor_sku'],product=l['name'],variation=l['variation'],pack_size=l['pack_size'],requested_packs=l['total_packs']) for l in snapshot['lines']])
 validate_projection(p);return p


def validate_projection(p):
 if set(p)!=PUBLIC or p['protocol']!=VERSION or p['currency']!='USD' or type(p['revision_number']) is not int or p['revision_number']<1:raise InvalidWorkbook('Unsupported explicit quotation protocol/currency')
 if not 1<=len(p['lines'])<=MAX_LINES:raise InvalidWorkbook('Quotation supports 1–5,000 lines')
 for l in p['lines']:
  if set(l)!=LINE or type(l['pack_size']) is not int or not 1<=l['pack_size']<=1000000 or type(l['requested_packs']) is not int or not 0<=l['requested_packs']<=4000000:raise InvalidWorkbook('Invalid protected request quantity or pack')
  if any(l[k] is not None and (not isinstance(l[k],str) or len(l[k])>2000) for k in LINE-{'pack_size','requested_packs'}):raise InvalidWorkbook('Invalid protected product label')
 if len({l['line_id'] for l in p['lines']})!=len(p['lines']):raise InvalidWorkbook('Duplicate protected request line')


def signature(p,key):return hmac.new(key,digest(p).encode(),hashlib.sha256).hexdigest()


def stable_zip(content):
 output=BytesIO()
 with ZipFile(BytesIO(content)) as source,ZipFile(output,'w',ZIP_DEFLATED) as target:
  for name in sorted(source.namelist()):
   raw=source.read(name)
   if name=='docProps/core.xml':
    root=ET.fromstring(raw)
    for tag in ['created','modified']:
     element=root.find('{http://purl.org/dc/terms/}'+tag)
     if element is not None:element.text='2000-01-01T00:00:00Z'
    raw=ET.tostring(root,encoding='utf-8',xml_declaration=True)
   info=ZipInfo(name,(2000,1,1,0,0,0));info.compress_type=ZIP_DEFLATED;target.writestr(info,raw)
 return output.getvalue()


def template_bytes(p,sig):
 validate_projection(p);wb=Workbook();ws=wb.active;ws.title='Quotation';ws.sheet_view.showGridLines=False;ws.freeze_panes='H12';ws.sheet_properties.pageSetUpPr.fitToPage=True
 rows={2:[p['purchaser']],3:['PURCHASING REQUEST · QUOTATION ONLY'],4:['Supplier',p['supplier']],5:['Request',p['request_id'],None,None,'Revision',p['revision_number'],'Currency',p['currency']],6:['Ship to',p['ship_to']],7:['Edit only amber cells. Enter whole packs, explicit 0 if unavailable; blank means unanswered.'],8:['Prices are per purchasing pack, USD, up to 4 decimal places. No formulas. Do not change product or pack fields.'],9:['Freight USD',None,'Tax USD',None,'Discount USD',None],10:['Enter freight, tax and discount amounts explicitly; 0 means none. Blank means unknown. This is not a purchase order.'],11:HEADERS}
 for n,values in rows.items():
  for col,value in enumerate(values,1):
   cell=ws.cell(n,col,value)
   if isinstance(value,str):cell.data_type='s'
 for n,l in enumerate(p['lines'],START):
  values=[l['line_id'],l['sku'],l['vendor_sku'],l['product'],l['variation'],l['pack_size'],l['requested_packs'],None,None]
  for col,value in enumerate(values,1):
   cell=ws.cell(n,col,value)
   if isinstance(value,str):cell.data_type='s'
   cell.font=Font(name='Arial',size=10,color='243047');cell.alignment=Alignment(vertical='center',wrap_text=col in (4,5),horizontal='right' if col>=6 else 'left')
  ws.row_dimensions[n].height=30
  for col in (8,9):ws.cell(n,col).protection=Protection(locked=False);ws.cell(n,col).fill=PatternFill('solid',fgColor='FFF0C2');ws.cell(n,col).number_format='0' if col==8 else '0.0000'
 for coord in ['B9','D9','F9']:ws[coord].protection=Protection(locked=False);ws[coord].fill=PatternFill('solid',fgColor='FFF0C2');ws[coord].number_format='0.00'
 for cell in ws[11]:cell.font=Font(name='Arial',size=10,bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='344B62');cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
 ws.merge_cells('B5:D5');ws.merge_cells('B6:I6');ws.merge_cells('A7:I7');ws.merge_cells('A8:I8');ws.merge_cells('A10:I10');ws.row_dimensions[6].height=min(409,max(34,14*((len(p['ship_to'])+149)//150+p['ship_to'].count('\n'))));ws['B6'].alignment=Alignment(wrap_text=True,vertical='center');ws.row_dimensions[11].height=34;ws['A2'].font=Font(name='Arial',size=16,bold=True,color='243047');ws['A3'].font=Font(name='Arial',size=12,bold=True,color='526C81')
 for col,width in zip('ABCDEFGHI',[38,18,18,30,30,13,14,17,21]):ws.column_dimensions[col].width=width
 ws.protection.sheet=True;ws.protection.selectLockedCells=True;ws.protection.selectUnlockedCells=False;ws.auto_filter.ref=f'A11:I{START+len(p["lines"])-1}'
 integer=DataValidation(type='whole',operator='between',formula1=0,formula2=4000000,allow_blank=True);integer.errorTitle='Whole purchasing packs';integer.error='Enter a nonnegative whole number of packs.';integer.showErrorMessage=True;ws.add_data_validation(integer);integer.add(f'H{START}:H{START+len(p["lines"])-1}')
 money=DataValidation(type='decimal',operator='between',formula1=0,formula2=9999999999.9999,allow_blank=True);money.showErrorMessage=True;ws.add_data_validation(money);money.add(f'I{START}:I{START+len(p["lines"])-1}');money.add('B9');money.add('D9');money.add('F9')
 manifest=wb.create_sheet('Protocol');manifest.sheet_view.showGridLines=False;manifest.column_dimensions['A'].width=24;manifest.column_dimensions['B'].width=80
 data=dict(protocol=VERSION,request_id=p['request_id'],vendor_id=p['vendor_id'],revision_id=p['revision_id'],revision_number=p['revision_number'],currency=p['currency'],snapshot_digest=digest(p),signature=sig)
 for n,k in enumerate(MANIFEST,1):manifest.cell(n,1,k);cell=manifest.cell(n,2,data[k]);cell.data_type='s' if isinstance(data[k],str) else 'n'
 manifest.protection.sheet=True;wb.properties.creator='ERUPTED';wb.properties.lastModifiedBy='ERUPTED';wb.properties.title='Purchasing quotation';output=BytesIO();wb.save(output);return stable_zip(output.getvalue())


def bounded_package(content):
 if not isinstance(content,bytes) or not content or len(content)>MAX_FILE:raise InvalidWorkbook('Upload one nonempty XLSX, at most 10 MiB')
 try:
  with ZipFile(BytesIO(content)) as z:
   entries=z.infolist();names=[e.filename for e in entries]
   if len(entries)>MAX_ENTRIES or len(set(names))!=len(names) or len({n.casefold() for n in names})!=len(names):raise InvalidWorkbook('Duplicate or excessive workbook ZIP entries')
   if sum(e.file_size for e in entries)>MAX_EXPANDED:raise InvalidWorkbook('Expanded workbook exceeds 64 MiB')
   allowed=re.compile(r'(\[Content_Types\]\.xml|_rels/\.rels|docProps/(app|core)\.xml|xl/(workbook\.xml|styles\.xml|sharedStrings\.xml|theme/theme1\.xml|_rels/workbook\.xml\.rels|worksheets/sheet[12]\.xml))\Z')
   for e in entries:
    if not allowed.fullmatch(e.filename) or e.flag_bits&1 or e.compress_type not in (0,8) or e.file_size>32*1024*1024 or (e.file_size>1024*1024 and e.file_size>max(e.compress_size,1)*200):raise InvalidWorkbook('Unsupported, encrypted, active or oversized workbook part: '+e.filename[:120])
    raw=z.read(e);xml=raw.decode('utf-8-sig');upper=xml.upper()
    if '<!DOCTYPE' in upper or '<!ENTITY' in upper or '\x00' in xml:raise InvalidWorkbook('DTD/entity declarations and unsupported XML encodings are prohibited')
    depth=0;nodes=0;addresses=set();row_ids=set();current_row=0
    for event,element in ET.iterparse(BytesIO(raw),events=('start','end')):
     if event=='start':
      depth+=1;nodes+=1
      if depth>32 or nodes>350000 or len(element.attrib)>32:raise InvalidWorkbook('Workbook XML exceeds safe structural bounds')
      tag=element.tag.split('}')[-1]
      if any(v in element.attrib.get('ContentType','').lower() for v in ('macroenabled','vba','oleobject')):raise InvalidWorkbook('Macro-enabled and active workbook content types are prohibited')
      if tag in ('f','externalLink','externalReference','oleObject','embeddedObject','hyperlink','legacyDrawing','extLst','formula'):raise InvalidWorkbook('Formulas, external links, embedded/active workbook objects are prohibited')
      if tag=='Relationship' and (element.attrib.get('TargetMode')=='External' or any(word in element.attrib.get('Type','').lower() for word in ('vba','externallink','oleobject'))):raise InvalidWorkbook('External or active workbook relationship prohibited')
      if tag=='definedName' and element.attrib.get('name') not in ('_xlnm.Print_Area','_xlnm.Print_Titles','_xlnm._FilterDatabase'):raise InvalidWorkbook('Unexpected workbook named range prohibited')
      if tag=='c':
       address=element.attrib.get('r','');match=re.fullmatch(r'[A-I]([1-9]\d{0,3})',address)
       if not match:raise InvalidWorkbook('Cell address exceeds the supported quotation layout')
       if address in addresses or int(match[1])!=current_row:raise InvalidWorkbook('Duplicate cell address or conflicting worksheet row identity')
       addresses.add(address)
      if tag=='row':
       current_row=int(element.attrib.get('r','0'))
       if not 1<=current_row<=MAX_LINES+START or current_row in row_ids:raise InvalidWorkbook('Duplicate or oversized workbook row identity')
       row_ids.add(current_row)
      if tag=='dimension':
       bounds=re.fullmatch(r'[A-I]([1-9]\d{0,3})(?::[A-I]([1-9]\d{0,3}))?',element.attrib.get('ref',''))
       if not bounds or any(int(n)>MAX_LINES+START for n in bounds.groups() if n):raise InvalidWorkbook('Workbook dimensions exceed the quotation resource bound')
     else:
      if element.tag.split('}')[-1] in ('formula1','formula2') and element.text and not re.fullmatch(r'\d+(?:\.\d+)?',element.text):raise InvalidWorkbook('Only literal numeric validation bounds are allowed')
      if element.tag.split('}')[-1]=='definedName' and element.text and not re.fullmatch(r"'?(Quotation|Protocol)'?!\$?[A-I]?\$?\d+(?::\$?[A-I]?\$?\d+)?",element.text):raise InvalidWorkbook('Named ranges must be bounded local quotation ranges')
      if element.text and len(element.text)>4096:raise InvalidWorkbook('Workbook text exceeds safe cell bound')
      depth-=1;element.clear()
   if '[Content_Types].xml' not in names or 'xl/workbook.xml' not in names:raise InvalidWorkbook('Not an XLSX workbook package')
 except InvalidWorkbook:raise
 except (BadZipFile,ValueError,UnicodeError,ET.ParseError,KeyError,RuntimeError,OverflowError) as e:raise InvalidWorkbook('Malformed or unsupported XLSX container: '+str(e)[:180]) from None


def number(value,*,scale,limit,label):
 if value is None:return None
 if isinstance(value,bool) or not isinstance(value,(str,int,float,Decimal)):raise InvalidWorkbook(label+': literal number required')
 raw=str(value)
 if not re.fullmatch(r'\d+(?:\.\d+)?',raw) or len(raw)>30:raise InvalidWorkbook(label+': nonnegative literal decimal required')
 try:d=Decimal(raw)
 except InvalidOperation:raise InvalidWorkbook(label+': invalid number') from None
 if not d.is_finite() or d>Decimal(limit) or d!=d.quantize(Decimal(1).scaleb(-scale)):raise InvalidWorkbook(label+f': at most {scale} decimal places; limit {limit}')
 return format(d.normalize(),'f')


def parse(content,p,sig):
 bounded_package(content);validate_projection(p)
 try:
  wb=load_workbook(BytesIO(content),data_only=False,read_only=False,keep_links=False)
  if wb.sheetnames!=['Quotation','Protocol'] or any(s.sheet_state!='visible' for s in wb) or wb.defined_names.keys()-{'_xlnm.Print_Area','_xlnm.Print_Titles','_xlnm._FilterDatabase'}:raise InvalidWorkbook('Expected exactly the visible Quotation and Protocol sheets')
  ws=wb['Quotation'];manifest=wb['Protocol']
  if ws.max_column>9 or ws.max_row>MAX_LINES+START or manifest.max_column!=2 or manifest.max_row!=8:raise InvalidWorkbook('Unexpected workbook dimensions')
  expected=dict(protocol=VERSION,request_id=p['request_id'],vendor_id=p['vendor_id'],revision_id=p['revision_id'],revision_number=p['revision_number'],currency=p['currency'],snapshot_digest=digest(p),signature=sig)
  actual={}
  for row in manifest.iter_rows(values_only=True):
   if row[0] in actual:raise InvalidWorkbook('Duplicate protocol field')
   actual[row[0]]=row[1]
  if isinstance(actual.get('revision_number'),bool):raise InvalidWorkbook('Request revision must be a literal integer')
  if actual!=expected:raise InvalidWorkbook('Request, vendor, revision, currency or signed request snapshot does not match')
  protected={'A4':'Supplier','A5':'Request','E5':'Revision','G5':'Currency','A6':'Ship to','A9':'Freight USD','C9':'Tax USD','E9':'Discount USD','A2':p['purchaser'],'A3':'PURCHASING REQUEST · QUOTATION ONLY','B4':p['supplier'],'B5':p['request_id'],'F5':p['revision_number'],'H5':p['currency'],'B6':p['ship_to']}
  for address,value in protected.items():
   if ws[address].value!=value or isinstance(value,int) and isinstance(ws[address].value,bool):raise InvalidWorkbook('Protected request information changed at '+address)
  if [c.value for c in ws[11]]!=HEADERS:raise InvalidWorkbook('Quotation column headings changed')
  by={l['line_id']:l for l in p['lines']};seen=set();lines=[];issues=[]
  for cells in ws.iter_rows(min_row=START,max_col=9):
   values=[c.value for c in cells]
   if all(v is None for v in values):continue
   token=values[0]
   if token not in by:raise InvalidWorkbook('Unexpected product/line identity at row '+str(cells[0].row))
   if token in seen:raise InvalidWorkbook('Duplicate product line: '+by[token]['product'])
   seen.add(token);l=by[token];wanted=[l['line_id'],l['sku'],l['vendor_sku'],l['product'],l['variation'],l['pack_size'],l['requested_packs']]
   if values[:7]!=wanted or any(isinstance(value,bool) for value in values[:7]):raise InvalidWorkbook(l['product']+': protected product, requested quantity or pack configuration changed')
   qty=number(values[7],scale=0,limit='4000000',label=l['product']+' available packs');price=number(values[8],scale=4,limit='9999999999.9999',label=l['product']+' price per pack')
   quantity=int(qty) if qty is not None else None
   status='MISSING' if quantity is None else 'UNAVAILABLE' if quantity==0 else 'PARTIAL' if quantity<l['requested_packs'] else 'EXCESS' if quantity>l['requested_packs'] else 'FULL'
   if quantity is None:issues.append(l['product']+': quantity unanswered')
   if quantity and price is None:issues.append(l['product']+': positive availability has no quoted price')
   lines.append(dict(line_id=token,available_packs=quantity,quoted_pack_price=price,status=status))
  if seen!=set(by):raise InvalidWorkbook('Missing request lines: '+', '.join(by[k]['product'] for k in set(by)-seen))
  components={k:number(ws[cell].value,scale=2,limit='999999999999.99',label=k+' amount') for k,cell in [('freight','B9'),('tax','D9'),('discount','F9')]}
  issues.extend(k+': amount unanswered' for k,v in components.items() if v is None)
  return dict(protocol=VERSION,request_id=p['request_id'],revision_id=p['revision_id'],currency=p['currency'],lines=sorted(lines,key=lambda r:r['line_id']),components=components,issues=sorted(issues))
 except InvalidWorkbook:raise
 except (ValueError,KeyError,TypeError,OverflowError,IndexError) as e:raise InvalidWorkbook('Unsupported workbook structure: '+str(e)[:180]) from None
 finally:
  if 'wb' in locals():wb.close()
