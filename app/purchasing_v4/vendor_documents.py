"""Vendor-safe immutable projections. Renderers cannot accept internal PO objects.

Runtime generation uses the application's existing pinned server-side libraries;
there is no Node service or external document transport in the deployment path.
"""
from io import BytesIO
from html import escape
from datetime import datetime,timezone
from zipfile import ZipFile,ZipInfo,ZIP_DEFLATED
import hashlib
import math
from xml.etree import ElementTree as ET
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.lib.enums import TA_RIGHT
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side

TOP={'schema','kind','number','issued_at','purchaser','supplier','ship_to','lines'}
LINE={'sku','vendor_sku','product','variation','pack_size','packs','units'}


def vendor_projection(child_id,po_id,kind,snapshot):
    result=dict(schema='v4-vendor-document/1',kind=kind,number='PO #'+str(po_id) if kind=='PO' else 'Request '+child_id,
      issued_at=snapshot['authorized_at'],purchaser=snapshot['header']['purchaser'],supplier=snapshot['supplier']['name'],ship_to=snapshot['header']['ship_to'],
      lines=[dict(sku=r['sku'],vendor_sku=r['vendor_sku'],product=r['name'],variation=r['variation'],pack_size=r['pack_size'],packs=r['total_packs'],units=r['total_units']) for r in snapshot['lines']])
    validate_projection(result);return result


def validate_projection(p):
    if not isinstance(p,dict) or set(p)!=TOP or p['schema']!='v4-vendor-document/1' or p['kind'] not in ('PO','REQUEST'):raise ValueError('Only the exact vendor-safe document projection is accepted')
    if not isinstance(p['lines'],list) or not 1<=len(p['lines'])<=5000:raise ValueError('Bounded nonempty vendor document required')
    for k in TOP-{'lines'}:
        if not isinstance(p[k],str) or not p[k] or len(p[k])>2000:raise ValueError('Invalid public header')
    datetime.fromisoformat(p['issued_at'])
    for r in p['lines']:
        if not isinstance(r,dict) or set(r)!=LINE:raise ValueError('Internal fields are prohibited in vendor documents')
        if any(not isinstance(r[k],str) or len(r[k])>2000 for k in ('product','variation')):raise ValueError('Invalid public product label')
        if any(r[k] is not None and (not isinstance(r[k],str) or len(r[k])>2000) for k in ('sku','vendor_sku')):raise ValueError('Invalid public identifier')
        if any(type(r[k]) is not int or r[k]<1 for k in ('pack_size','packs','units')) or r['units']!=r['pack_size']*r['packs']:raise ValueError('Exact public pack/unit conversion required')


def pdf_bytes(p):
    validate_projection(p);output=BytesIO();styles=getSampleStyleSheet()
    styles.add(ParagraphStyle('CellV4',fontName='Helvetica',fontSize=8,leading=11,textColor=colors.HexColor('#28354b')))
    styles.add(ParagraphStyle('NumV4',parent=styles['CellV4'],alignment=TA_RIGHT))
    def para(s,style='CellV4'):return Paragraph(escape(str(s or '')).replace('\n','<br/>'),styles[style])
    title='PURCHASE ORDER' if p['kind']=='PO' else 'PURCHASING REQUEST — NOT A PURCHASE ORDER'
    flow=[Paragraph(escape(p['purchaser']),styles['Title']),Spacer(1,12),Paragraph(escape(title),styles['Heading2']),para(p['number']),Spacer(1,14),para('Supplier: '+p['supplier']),para('Issued: '+p['issued_at'][:10]),Spacer(1,10),para('Ship to: '+p['ship_to']),Spacer(1,18)]
    table=[[para(v) for v in ['Our SKU / Vendor SKU','Product / Variation','Units per pack','Packs','Total units']]]
    for r in p['lines']:
        table.append([para((r['sku'] or '—')+'\n'+(r['vendor_sku'] or '—')),para(r['product']+'\n'+r['variation']),para(r['pack_size'],'NumV4'),para(r['packs'],'NumV4'),para(r['units'],'NumV4')])
    table.append([para('TOTAL'),para(''),para(''),para(sum(r['packs'] for r in p['lines']),'NumV4'),para(sum(r['units'] for r in p['lines']),'NumV4')])
    t=Table(table,colWidths=[100,220,65,55,70],repeatRows=1,hAlign='LEFT');t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e7edf2')),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),8),('RIGHTPADDING',(0,0),(-1,-1),8),('TOPPADDING',(0,0),(-1,-1),8),('BOTTOMPADDING',(0,0),(-1,-1),8),('LINEBELOW',(0,0),(-1,-1),.4,colors.HexColor('#dfe5ed')),('BACKGROUND',(0,-1),(-1,-1),colors.HexColor('#edf3f2'))]));flow.append(t)
    flow.extend([Spacer(1,15),para('Quantities are consolidated across the order. Pack sizes and quantities are fixed in this issued snapshot.'),Spacer(1,6),para('ISOLATED DEVELOPMENT SAMPLE — NOT SENT TO A VENDOR')])
    def page(canvas,doc):
        canvas.setFont('Helvetica',8);canvas.setFillColor(colors.HexColor('#718097'));canvas.drawString(42,28,p['number']);canvas.drawRightString(570,28,'Page '+str(doc.page))
    SimpleDocTemplate(output,pagesize=(612,792),leftMargin=42,rightMargin=42,topMargin=40,bottomMargin=48,title=p['number'],author='ERUPTED',subject='Vendor purchasing document',creator='ERUPTED Purchasing',invariant=1).build(flow,onFirstPage=page,onLaterPages=page)
    return output.getvalue()


def xlsx_bytes(p):
    validate_projection(p);book=Workbook();sheet=book.active;sheet.title='Purchase Order' if p['kind']=='PO' else 'Purchasing Request'
    book.properties.creator='ERUPTED';book.properties.lastModifiedBy='ERUPTED';book.properties.title=p['number'];book.properties.subject='Vendor purchasing document';book.properties.description='Immutable consolidated quantities';book.properties.created=datetime(2000,1,1);book.properties.modified=datetime(2000,1,1)
    def text(row,col,value):
        cell=sheet.cell(row,col);cell.value=str(value or '');cell.data_type='s';return cell
    for row,label,value in [(1,p['purchaser'],''),(2,'PURCHASE ORDER' if p['kind']=='PO' else 'PURCHASING REQUEST — NOT A PURCHASE ORDER',p['number']),(3,'Supplier',p['supplier']),(4,'Issued',p['issued_at'][:10]),(5,'Ship to',p['ship_to']),(6,'ISOLATED DEVELOPMENT SAMPLE — NOT SENT TO A VENDOR','')]:
        text(row,1,label);text(row,2,value)
    headers=['Our SKU','Vendor SKU','Product','Variation','Units per pack','Packs','Total units']
    for i,h in enumerate(headers,1):text(8,i,h)
    for n,r in enumerate(p['lines'],9):
        for i,k in enumerate(['sku','vendor_sku','product','variation'],1):text(n,i,r[k])
        for i,k in enumerate(['pack_size','packs','units'],5):sheet.cell(n,i,r[k]).number_format='#,##0'
    last=9+len(p['lines']);text(last,1,'TOTAL');sheet.cell(last,6,sum(r['packs'] for r in p['lines']));sheet.cell(last,7,sum(r['units'] for r in p['lines']))
    for row in sheet:
        for cell in row:
            cell.font=Font(name='Arial',size=10,color='26364A');cell.alignment=Alignment(vertical='top',wrap_text=True,horizontal='right' if cell.column>=5 else 'left')
    sheet['A1'].font=Font(name='Arial',size=19,bold=True,color='24364B');sheet['A2'].font=Font(name='Arial',size=11,bold=True,color='50677A')
    for cell in sheet[8]:cell.font=Font(name='Arial',size=10,bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='344C62')
    for cell in sheet[last]:cell.fill=PatternFill('solid',fgColor='E7F0ED');cell.font=Font(name='Arial',size=10,bold=True)
    for col,width in zip('ABCDEFG',[18,20,34,34,16,12,16]):sheet.column_dimensions[col].width=width
    sheet['B2']=None;text(2,5,p['number']);sheet.merge_cells('A2:D2');sheet.merge_cells('E2:G2');sheet.row_dimensions[2].height=40;sheet.merge_cells('B3:G3');sheet.row_dimensions[3].height=min(409,max(24,math.ceil(len(p['supplier'])/100)*14));sheet.merge_cells('A6:G6');sheet.row_dimensions[6].height=22
    sheet.merge_cells('A1:G1');sheet.merge_cells('B5:G5');sheet.row_dimensions[1].height=min(409,max(32,math.ceil(len(p['purchaser'])/65)*24));sheet.row_dimensions[5].height=min(409,max(42,(math.ceil(len(p['ship_to'])/110)+p['ship_to'].count('\n'))*14));sheet.row_dimensions[8].height=26
    for row in range(9,last):sheet.row_dimensions[row].height=min(409,max(34,14*max(math.ceil(len(str(sheet.cell(row,col).value or ''))/width) for col,width in [(1,18),(2,20),(3,34),(4,34)])))
    sheet.freeze_panes='C9';sheet.auto_filter.ref=f'A8:G{last-1}';sheet.sheet_view.showGridLines=False;sheet.print_title_rows='1:8';sheet.print_area=f'A1:G{last}';sheet.sheet_properties.pageSetUpPr.fitToPage=True;sheet.page_setup.orientation='landscape';sheet.page_setup.paperSize=sheet.PAPERSIZE_A4;sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0
    output=BytesIO();book.save(output)
    # Reproducible ZIP metadata: retries cannot quietly issue different originals.
    normalized=BytesIO()
    with ZipFile(BytesIO(output.getvalue())) as source,ZipFile(normalized,'w',ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            info=ZipInfo(name,(2000,1,1,0,0,0));info.compress_type=ZIP_DEFLATED;content=source.read(name)
            if name=='docProps/core.xml':
                root=ET.fromstring(content)
                for tag in ['created','modified']:
                    element=root.find('{http://purl.org/dc/terms/}'+tag)
                    if element is not None:element.text='2000-01-01T00:00:00Z'
                content=ET.tostring(root,encoding='utf-8',xml_declaration=True)
            target.writestr(info,content)
    return normalized.getvalue()


def render(p):return {'PDF':pdf_bytes(p),'XLSX':xlsx_bytes(p)}
