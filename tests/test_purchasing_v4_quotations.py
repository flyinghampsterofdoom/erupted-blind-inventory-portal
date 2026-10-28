"""Disposable PostgreSQL quotation protocol and accepted Phase5 authorization."""
from io import BytesIO
from uuid import uuid4
from copy import deepcopy
from decimal import Decimal
from concurrent.futures import ThreadPoolExecutor
from zipfile import ZipFile,ZIP_DEFLATED
import hashlib,json
import pytest
from sqlalchemy import text
from openpyxl import load_workbook
from app.purchasing_v4 import quotations as q,submission,controls,outbox,worksheets
from app.purchasing_v4 import quotation_workbooks as xlsx
from app.purchasing_v4.quotation_finance import CONTRACT
from app.purchasing_v4.schema import migrate_quotations
from test_purchasing_v4_submission import purchasing,sheet,ready,engine,template,reading,review,authorize,fresh,edit,save,pytestmark
NS='synthetic-fixture-source'

@pytest.fixture
def quote(purchasing):
 e,state=purchasing;sid=state['rows'][0]['supplier_id']
 with e.begin() as c:
  migrate_quotations(c);h=controls.heads(c,{('SUPPLIER',sid)})[('SUPPLIER',sid)]
  controls.apply(c,actor_id=1,command_id=str(uuid4()),reason='Explicit synthetic confirmation workflow',changes=[dict(kind='SUPPLIER',key=sid,expected_version=h['version'],value=dict(h['payload'],confirmation_enabled=True))])
 result=authorize(purchasing,review(purchasing));request=result['children'][0]['id']
 with e.begin() as c:head=q.prepare(c,actor_id=1,command_id=str(uuid4()),request_id=request)
 return e,state,request,head


def book(quote,quantity=6,price='50.1234',components=('0','0','0'),mutate=None):
 e,state,rid,head=quote
 with e.connect() as c:rev=q.revision(c,head['revision_id']);content=bytes(rev['template'])
 wb=load_workbook(BytesIO(content));ws=wb['Quotation'];ws['H12']=quantity;ws['I12']=price
 for cell,value in zip(['B9','D9','F9'],components):ws[cell]=value
 if mutate:mutate(wb)
 out=BytesIO();wb.save(out);wb.close();return out.getvalue()


def upload(quote,content=None,**options):
 return q.intake(quote[0],actor_id=1,request_id=quote[2],filename='vendor.xlsx',content=content or book(quote,**options))


def select(quote,response,version=None):
 e,_,rid,head=quote
 with e.begin() as c:current=q.head(c,rid);return q.transition(c,actor_id=1,command_id=str(uuid4()),request_id=rid,expected_version=current['version'] if version is None else version,kind='SELECT_RESPONSE',reason='Owner explicitly selects this retained vendor response',response_id=response['response_id'])


def final(quote,*,response=None,choices=None,components=None,**options):
 e,state,rid,head=quote
 if response is None:response=upload(quote);select(quote,response)
 with e.connect() as c:current=q.head(c,rid);snap=q.revision(c,current['revision_id'])['snapshot']
 values={l['row_id']:dict(packs=l['packs'],price_action='ACCEPT_QUOTE',pack_price=None,reason='Owner accepts exact vendor pack quotation',quantity_reason='',excess_reason='') for l in snap['lines']}
 args=dict(actor_id=1,review_id=str(uuid4()),request_id=rid,expected_version=current['version'],response_id=response['response_id'],choices=choices or values,components=components or {k:dict(source='VENDOR_QUOTE',amount=None,reason='Owner accepts explicitly quoted amount') for k in ['freight','tax','discount']},contract=CONTRACT,unit_adapter_approved=True,release_remainder=True,incremental_reason='Explicit additional synthetic units after reviewing all current supply and pending requests',source_namespace=NS)
 args.update(options);return q.prepare_final_review(e,**args)


def commit(quote,r,cid=None):
 with quote[0].begin() as c:return submission.authorize(c,actor_id=1,command_id=cid or str(uuid4()),review_id=r['id'],acknowledged=True,source_namespace=NS)


def test_valid_import_is_evidence_only_and_explicit_final_uses_phase5_intents(quote):
 e,state,rid,h=quote
 from app.purchasing_v4.reconciliation import public_fingerprints
 with reading(e) as c:before_public=public_fingerprints(c)
 with e.connect() as c:before=c.execute(text('SELECT count(*) FROM public.purchase_orders')).scalar_one();intents=set(c.execute(text('SELECT id::text FROM purchasing_v4.purchasing_intent WHERE vendor_authorization_id=:id'),{'id':rid}).scalars())
 response=upload(quote)
 with reading(e) as c:assert public_fingerprints(c)==before_public
 with e.connect() as c:
  assert c.execute(text('SELECT count(*) FROM public.purchase_orders')).scalar_one()==before
  assert q.head(c,rid)['selected_response_id'] is None
 select(quote,response);r=final(quote,response=response);assert not r['blocking'];result=commit(quote,r);assert result['commitments_created']==1
 with e.connect() as c:
  assert set(c.execute(text('SELECT intent_id::text FROM purchasing_v4.intent_consumption')).scalars())==intents
  assert q.head(c,rid)['state']=='COMMITTED'
  assert worksheets.snapshot(c,state['session']['id'],1)['rows'][0]['sealed']['resolved_purchase_order_id']==result['children'][0]['purchase_order_id']
  assert type(worksheets.snapshot(c,state['session']['id'],1)['rows'][0]['sealed']['approved_quote_units']) is int
  assert c.execute(text('SELECT count(*) FROM purchasing_v4.active_request_reservation')).scalar_one()==0
  assert c.execute(text('SELECT count(*) FROM purchasing_v4.purchasing_intent')).scalar_one()==4
  assert c.execute(text('SELECT count(*) FROM purchasing_v4.quotation_financial')).scalar_one()==1
 assert commit(quote,r)['authorization_id']==result['authorization_id']


def test_exact_decimal_charges_and_legacy_finance_total(quote):
 response=upload(quote,components=('1.25','0.03','0.02'));select(quote,response);r=final(quote,response=response);assert not r['blocking'];f=r['quotation']['financial']
 assert f['merchandise']=='300.74' and f['total']=='302.00';assert r['lines'][0]['unit_adapter_delta'] is not None
 result=commit(quote,r)
 from sqlalchemy.orm import Session
 from app.services.v2_order_payments_service import _order_cost_snapshot
 with Session(quote[0]) as db:
  assert _order_cost_snapshot(db,result['children'][0]['purchase_order_id'])==(Decimal('302.00'),True)
  from app.services.purchase_order_admin_service import list_purchase_orders
  assert next(x for x in list_purchase_orders(db) if x['id']==result['children'][0]['purchase_order_id'])['invoice_amount']==Decimal('302.00')


def test_raw_and_semantic_duplicate_uploads_are_idempotent(quote):
 content=book(quote);first=upload(quote,content);assert upload(quote,content)['id']==first['id']
 other=book(quote,price='50.12340000');second=upload(quote,other);assert second['response_id']==first['response_id']
 with quote[0].connect() as c:assert c.execute(text('SELECT count(*) FROM purchasing_v4.quotation_response')).scalar_one()==1


def test_unknown_and_zero_prices_remain_distinct_and_require_explicit_authority(quote):
 unknown=upload(quote,price=None);select(quote,unknown);r=final(quote,response=unknown);assert r['blocking'];assert r['committed_total'] is None
 with pytest.raises(controls.InvalidPolicy):commit(quote,r)
 zero=upload(quote,price='0');select(quote,zero);r=final(quote,response=zero);assert not r['blocking'] and r['committed_total']=='0.00';assert r['lines'][0]['reported']['quoted_pack_price']=='0'
 assert commit(quote,r)['commitments_created']==1


@pytest.mark.parametrize('mutation,message',[
 (lambda w:w['Quotation'].__setitem__('F12',19),'protected'),
 (lambda w:w['Quotation'].__setitem__('H12','=1+1'),'Formulas'),
 (lambda w:w['Quotation'].__setitem__('H12','1.5'),'decimal'),
 (lambda w:w['Quotation'].__setitem__('I12','-1'),'nonnegative'),
 (lambda w:w['Quotation'].__setitem__('I12','1.23456'),'decimal'),
 (lambda w:w['Protocol'].__setitem__('B6','EUR'),'match'),
 (lambda w:w['Quotation'].__setitem__('A9','Different amount basis'),'Protected'),
 (lambda w:w['Quotation'].delete_rows(12),'Missing'),
 (lambda w:w['Quotation'].append([w['Quotation']['A12'].value,*[w['Quotation'].cell(12,c).value for c in range(2,10)]]),'Duplicate'),
 (lambda w:w.create_sheet('Secret'),'part'),
])
def test_invalid_workbooks_are_quarantined_without_any_commitment(quote,mutation,message):
 e=quote[0]
 with reading(e) as c:
  from app.purchasing_v4.reconciliation import public_fingerprints
  before=public_fingerprints(c)
 result=upload(quote,book(quote,mutate=mutation));assert not result['validation']['valid'];assert message.lower() in str(result['validation']['errors']).lower()
 with reading(e) as c:
  assert public_fingerprints(c)==before
  assert c.execute(text('SELECT content FROM purchasing_v4.quotation_upload WHERE id=:id'),{'id':result['id']}).scalar_one()
  assert c.execute(text('SELECT count(*) FROM purchasing_v4.quotation_response')).scalar_one()==0


def test_missing_and_partial_quantities_require_explicit_allocations(quote):
 response=upload(quote,quantity=None);select(quote,response);r=final(quote,response=response);assert any('quantity exception' in s for s in r['blocking'])
 partial=upload(quote,quantity=3);select(quote,partial);r=final(quote,response=partial);assert r['blocking']
 choices=deepcopy(r['quotation']['owner_decisions']);line=next(iter(choices.values()));line['packs']=[1,1,1,0]
 r=final(quote,response=partial,choices=choices);assert not r['blocking'];assert r['lines'][0]['reported']['status']=='PARTIAL'
 assert commit(quote,r)['commitments_created']==1
 with quote[0].connect() as c:assert c.execute(text("SELECT sum((evidence->>'released_units')::bigint) FROM purchasing_v4.quotation_intent_resolution")).scalar_one()>0


def test_excess_is_preserved_and_only_explicit_revised_authorization_can_accept_it(quote):
 response=upload(quote,quantity=8);select(quote,response);r=final(quote,response=response);assert r['lines'][0]['reported']['available_packs']==8 and r['lines'][0]['reported']['status']=='EXCESS'
 choices=deepcopy(r['quotation']['owner_decisions']);line=next(iter(choices.values()));line['packs']=[2,2,3,1]
 blocked=final(quote,response=response,choices=choices);assert any('excess' in v for v in blocked['blocking'])
 line['excess_reason']='Owner explicitly authorizes the separately reviewed excess packs';approved=final(quote,response=response,choices=choices);assert not approved['blocking'];commit(quote,approved)
 with quote[0].connect() as c:assert c.execute(text("SELECT sum((evidence->>'additional_units')::bigint) FROM purchasing_v4.quotation_intent_resolution")).scalar_one()>0


def test_outdated_response_and_cancellation_invalidate_ready_review(quote):
 e,_,rid,h=quote;response=upload(quote);select(quote,response);r=final(quote,response=response)
 with e.begin() as c:
  current=q.head(c,rid);snapshot=q.revision(c,current['revision_id'])['snapshot'];changes={l['row_id']:dict(pack_size=l['pack_size'],packs=l['packs']) for l in snapshot['lines']}
  next_head=q.transition(c,actor_id=1,command_id=str(uuid4()),request_id=rid,expected_version=current['version'],kind='REISSUE',reason='Owner requests corrected quotation',changes=changes)
 with pytest.raises(controls.Conflict):commit(quote,r)
 with pytest.raises(controls.Conflict):select(quote,response)
 stale=upload(quote,book(quote,price='55'));assert stale['validation']['valid'] and stale['validation']['stale']
 with e.begin() as c:q.transition(c,actor_id=1,command_id=str(uuid4()),request_id=rid,expected_version=next_head['version'],kind='CANCEL',reason='Owner explicitly cancels the unresolved request')
 with e.connect() as c:
  assert c.execute(text('SELECT count(*) FROM purchasing_v4.active_request_reservation')).scalar_one()==0
  assert c.execute(text('SELECT count(*) FROM purchasing_v4.intent_consumption')).scalar_one()==0
  assert c.execute(text('SELECT count(*) FROM purchasing_v4.quotation_revision')).scalar_one()==2


def test_changed_active_response_and_current_supply_reject_stale_authorization(quote):
 first=upload(quote);select(quote,first);r=final(quote,response=first);second=upload(quote,price='60');select(quote,second)
 with pytest.raises(controls.Conflict):commit(quote,r)
 r=final(quote,response=second)
 with quote[0].begin() as c:c.execute(text('UPDATE public.purchase_order_store_allocations SET store_received_qty=store_received_qty+1 WHERE id=1'));c.execute(text('UPDATE public.purchase_order_lines SET received_qty_total=received_qty_total+1 WHERE id=1'))
 with pytest.raises(controls.Conflict):commit(quote,r)


def test_concurrent_review_authorization_consumes_original_intents_once(quote):
 response=upload(quote);select(quote,response);reviews=[final(quote,response=response),final(quote,response=response)]
 def run(r):
  try:return commit(quote,r)
  except controls.Conflict:return None
 with ThreadPoolExecutor(max_workers=2) as pool:results=list(pool.map(run,reviews))
 assert sum(r is not None for r in results)==1
 with quote[0].connect() as c:
  assert c.execute(text('SELECT count(*) FROM purchasing_v4.quotation_financial')).scalar_one()==1
  assert c.execute(text('SELECT count(*) FROM purchasing_v4.intent_consumption')).scalar_one()==4


def test_owner_candidate_missing_components_rejection_and_exception_are_explicit(quote):
 response=upload(quote,components=(None,None,None));select(quote,response);r=final(quote,response=response);assert r['committed_total'] is None
 choices=deepcopy(r['quotation']['owner_decisions']);next(iter(choices.values()))['price_action']='REJECT';components={k:dict(source='OWNER_EXCEPTION',amount='0',reason='Owner explicitly confirms no charge for this merchandise commitment') for k in ['freight','tax','discount']}
 r=final(quote,response=response,choices=choices,components=components);assert r['blocking']
 line=next(iter(choices.values()));line.update(price_action='OWNER_EXCEPTION',pack_price='45.5000',reason='Owner explicitly authorizes this alternative pack valuation; vendor quote retained')
 r=final(quote,response=response,choices=choices,components=components);assert not r['blocking'];commit(quote,r)


def test_interrupted_authorization_preserves_request_and_can_retry_forward(quote,monkeypatch):
 e=quote[0];r=final(quote);cid=str(uuid4());original=q.resolve
 def fail(*args,**kwargs):original(*args,**kwargs);raise RuntimeError('Injected commit interruption')
 monkeypatch.setattr(q,'resolve',fail)
 with pytest.raises(RuntimeError):commit(quote,r,cid)
 with e.connect() as c:
  assert q.head(c,quote[2])['state']=='OPEN'
  assert c.execute(text('SELECT count(*) FROM purchasing_v4.quotation_financial')).scalar_one()==0
  assert c.execute(text('SELECT count(*) FROM purchasing_v4.intent_consumption')).scalar_one()==0
 monkeypatch.setattr(q,'resolve',original);assert commit(quote,r,cid)['commitments_created']==1


def test_quotation_evidence_facts_immutable_and_worker_has_no_access(quote):
 e=quote[0];response=upload(quote);select(quote,response)
 for table in ['quotation_revision','quotation_upload','quotation_response','quotation_event','quotation_key']:
  with pytest.raises(Exception):
   with e.begin() as c:c.execute(text('DELETE FROM purchasing_v4.'+table))
 worker=outbox.worker_engine(e.url)
 for table in ['quotation_key','quotation_upload','quotation_response','quotation_financial']:
  with pytest.raises(Exception):
   with outbox.transaction(worker) as c:c.execute(text('SELECT * FROM purchasing_v4.'+table))
 worker.dispose()


@pytest.mark.parametrize('kind',['garbage','entity','external','macro','duplicate_zip','huge_dimension','duplicate_cell','row_mismatch','dimension'])
def test_malicious_containers_are_quarantined(quote,kind):
 original=book(quote)
 if kind=='garbage':content=b'not an XLSX file'
 else:
  result=BytesIO()
  with ZipFile(BytesIO(original)) as src,ZipFile(result,'w',ZIP_DEFLATED) as dst:
   for name in src.namelist():
    raw=src.read(name)
    if name=='xl/worksheets/sheet1.xml':
     if kind=='entity':raw=b'<!DOCTYPE x [<!ENTITY a "boom">]>'+raw
     if kind=='huge_dimension':raw=raw.replace(b'r="H12"',b'r="XFD99999999"')
     if kind=='duplicate_cell':
      import re
      cell=re.search(br'<c r="A2".*?</c>',raw)[0];raw=raw.replace(b'</row>',cell+b'</row>',1)
     if kind=='row_mismatch':raw=raw.replace(b'r="H12"',b'r="H13"')
     if kind=='dimension':
      import re
      raw=re.sub(br'<dimension[^>]*/>',b'<dimension ref="A1:XFD1048576"/>',raw)
    if kind=='external' and name=='xl/_rels/workbook.xml.rels':raw=raw.replace(b'</Relationships>',b'<Relationship Type="externalLink" Target="https://evil.invalid/" TargetMode="External" Id="attack"/></Relationships>')
    if kind=='macro' and name=='[Content_Types].xml':raw=raw.replace(b'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml',b'application/vnd.ms-excel.sheet.macroEnabled.main+xml')
    dst.writestr(name,raw)
    if kind=='duplicate_zip' and name=='xl/workbook.xml':dst.writestr(name,raw)
  content=result.getvalue()
 response=upload(quote,content);assert not response['validation']['valid']


def test_reissue_changes_pending_evidence_without_mutating_original_intent(quote):
 e,_,rid,h=quote
 with e.begin() as c:
  before=c.execute(text('SELECT sum(units) FROM purchasing_v4.purchasing_intent')).scalar_one();snapshot=q.revision(c,h['revision_id'])['snapshot'];changes={l['row_id']:dict(pack_size=10,packs=[1,0,0,0]) for l in snapshot['lines']}
  command=str(uuid4());args=dict(actor_id=1,command_id=command,request_id=rid,expected_version=h['version'],kind='REISSUE',reason='Owner explicitly requests a different pack configuration and quantity',changes=changes);first=q.transition(c,**args)
 with e.begin() as c:assert q.transition(c,**args)['revision_id']==first['revision_id']
 with e.connect() as c:
  assert c.execute(text('SELECT sum(units) FROM purchasing_v4.active_request_reservation')).scalar_one()==10
  assert c.execute(text('SELECT sum(units) FROM purchasing_v4.purchasing_intent')).scalar_one()==before
  assert q.revision(c,h['revision_id'])['snapshot']==snapshot


def test_request_consumption_cannot_bypass_quotation_authorization_at_database_boundary(quote):
 e,_,rid,_=quote
 with pytest.raises(Exception,match='canonical quoted commitment|quotation authorization'):
  with e.begin() as c:c.execute(text('INSERT INTO purchasing_v4.intent_consumption(intent_id,vendor_authorization_id) SELECT id,:request FROM purchasing_v4.purchasing_intent WHERE vendor_authorization_id=:request'),{'request':rid})
 with e.connect() as c:assert c.execute(text('SELECT count(*) FROM purchasing_v4.intent_consumption')).scalar_one()==0


def test_per_review_financial_candidate_and_unit_adapter_never_auto_approve(quote):
 response=upload(quote);select(quote,response)
 r=final(quote,response=response,contract=None,unit_adapter_approved=False,release_remainder=False)
 assert len(r['blocking'])>=3
 with pytest.raises(controls.InvalidPolicy):commit(quote,r)


def test_line_half_up_rounding_precedes_document_sum_and_zero_is_explicit(quote):
 from app.purchasing_v4.quotation_finance import calculate
 with quote[0].connect() as c:snapshot=deepcopy(q.revision(c,quote[3]['revision_id'])['snapshot'])
 first=snapshot['lines'][0];first.update(packs=[1,0,0,0],total_packs=1,units=[5,0,0,0],total_units=5,pack_size=5);second=deepcopy(first);second['row_id']=str(uuid4());snapshot['lines']=[first,second]
 response=dict(lines=[dict(line_id=l['row_id'],available_packs=1,quoted_pack_price='0.005',status='FULL') for l in snapshot['lines']],components={k:'0' for k in ['freight','tax','discount']})
 choices={l['row_id']:dict(packs=[1,0,0,0],price_action='ACCEPT_QUOTE',pack_price=None,reason='Owner explicitly approves this fractional-cent pack quotation',quantity_reason='',excess_reason='') for l in snapshot['lines']};components={k:dict(source='VENDOR_QUOTE',amount=None,reason='Explicitly quoted zero charge') for k in response['components']}
 lines,f,blocking=calculate(snapshot,response,choices,components,contract=CONTRACT,unit_adapter_approved=True,release_remainder=True)
 assert not blocking and [l['line_total'] for l in lines]==['0.01','0.01'] and f['total']=='0.02'


def test_quotation_template_allowlist_and_original_bytes_survive_future_configuration(quote):
 e,state,rid,h=quote
 with e.connect() as c:r=q.revision(c,h['revision_id']);original=bytes(r['template']);projection=r['public_projection']
 assert set(projection)==xlsx.PUBLIC and set(projection['lines'][0])==xlsx.LINE
 for key in ['expected_unit_cost','store_id','financing','valuation','Andresen','HWY 99','Longview']:
  with ZipFile(BytesIO(original)) as z:assert key.lower() not in '\n'.join(z.read(n).decode('utf8') for n in z.namelist()).lower()
 assert xlsx.template_bytes(projection,r['signature'])==original
 with e.begin() as c:
  sid=state['rows'][0]['supplier_id'];head=controls.heads(c,{('SUPPLIER',sid)})[('SUPPLIER',sid)];controls.apply(c,actor_id=1,command_id=str(uuid4()),reason='Future supplier name; old quotation stays immutable',changes=[dict(kind='SUPPLIER',key=sid,expected_version=head['version'],value=dict(head['payload'],name='Future supplier display'))])
 with e.connect() as c:assert bytes(q.revision(c,h['revision_id'])['template'])==original


def test_canceled_request_blocks_new_delivery_but_reconciles_prior_acceptance(quote):
 e,state,_,_=quote;sid=state['rows'][0]['supplier_id']
 with e.begin() as c:
  h=controls.heads(c,{('SUPPLIER',sid)})[('SUPPLIER',sid)];controls.apply(c,actor_id=1,command_id=str(uuid4()),reason='Explicit synthetic opted-in request delivery',changes=[dict(kind='SUPPLIER',key=sid,expected_version=h['version'],value=dict(h['payload'],automatic_spreadsheet_email=True,email='fixture@example.invalid'))])
 new=fresh((e,state));result=authorize(new,review(new));child=result['children'][0]
 with e.begin() as c:head=q.prepare(c,actor_id=1,command_id=str(uuid4()),request_id=child['id'])
 worker=outbox.worker_engine(e.url);outbox.generate_documents(worker,child['document_command_id'])
 with worker.connect() as c:delivery=str(c.execute(text("SELECT id FROM purchasing_v4.integration_command WHERE vendor_authorization_id=:id AND kind='DELIVER_DOCUMENTS'"),{'id':child['id']}).scalar_one())
 fake=outbox.SandboxProvider();fake.lose_ack=True;assert outbox.deliver_sandbox(worker,delivery,fake)['state']=='AMBIGUOUS'
 with e.begin() as c:q.transition(c,actor_id=1,command_id=str(uuid4()),request_id=child['id'],expected_version=head['version'],kind='CANCEL',reason='Owner cancels pending request; prior delivery acceptance still requires reconciliation')
 with pytest.raises(ValueError):outbox.deliver_sandbox(worker,delivery,outbox.SandboxProvider(),reconcile=True)
 fake.lose_ack=False;assert outbox.deliver_sandbox(worker,delivery,fake,reconcile=True)['state']=='ACKNOWLEDGED';assert len(fake.accepted)==1
 worker.dispose()


def test_request_to_quoted_commitment_link_is_unique_and_immutable(quote):
 r=final(quote);result=commit(quote,r);child=result['children'][0]['id']
 with pytest.raises(Exception):
  with quote[0].begin() as c:c.execute(text('INSERT INTO purchasing_v4.quotation_commitment(request_id,vendor_authorization_id) VALUES (:r,:v)'),dict(r=quote[2],v=child))
 with pytest.raises(Exception):
  with quote[0].begin() as c:c.execute(text('DELETE FROM purchasing_v4.quotation_commitment'))
 with quote[0].connect() as c:assert q.details(c,actor_id=1,request_id=quote[2])['committed_result']['authorization_id']==result['authorization_id']


@pytest.mark.parametrize('attack',['traversal','oversize','zip_count','compression_ratio','xml_depth'])
def test_workbook_resource_and_path_limits_fail_closed(quote,attack):
 if attack=='oversize':
  with pytest.raises(controls.InvalidPolicy,match='10 MiB'):q.intake(quote[0],actor_id=1,request_id=quote[2],filename='large.xlsx',content=b'x'*(xlsx.MAX_FILE+1))
  return
 original=book(quote);result=BytesIO()
 with ZipFile(BytesIO(original)) as src,ZipFile(result,'w',ZIP_DEFLATED) as dst:
  for name in src.namelist():
   raw=src.read(name)
   if attack=='xml_depth' and name=='xl/worksheets/sheet1.xml':raw=b'<x>'*34+raw+b'</x>'*34
   dst.writestr(name,raw)
  if attack=='traversal':dst.writestr('../outside.txt',b'not extracted')
  if attack=='zip_count':
   for i in range(129):dst.writestr('extra'+str(i),b'x')
  if attack=='compression_ratio':dst.writestr('xl/sharedStrings.xml',b'a'*(2*1024*1024))
 response=upload(quote,result.getvalue());assert response['validation']['valid'] is False


def test_boolean_cannot_impersonate_numeric_request_revision(quote):
 def mutate(w):w['Protocol']['B5']=True;w['Quotation']['F5']=True
 response=upload(quote,book(quote,mutate=mutate));assert response['validation']['valid'] is False


def test_direct_order_and_quotation_share_supply_guards_and_allow_reviewed_increment(quote):
 e,state,rid,_=quote;response=upload(quote);select(quote,response);stale=final(quote,response=response);sid=state['rows'][0]['supplier_id']
 with e.begin() as c:
  h=controls.heads(c,{('SUPPLIER',sid)})[('SUPPLIER',sid)];controls.apply(c,actor_id=1,command_id=str(uuid4()),reason='Explicit synthetic direct-order configuration for a separate owner-reviewed incremental purchase',changes=[dict(kind='SUPPLIER',key=sid,expected_version=h['version'],value=dict(h['payload'],confirmation_enabled=False))])
 direct=fresh((e,state));po=authorize(direct,review(direct))
 with pytest.raises(controls.Conflict):commit(quote,stale)
 current=final(quote,response=response);assert not current['blocking'];assert any(v['reference'].startswith('po:'+str(po['children'][0]['purchase_order_id'])+'/') for overlaps in current['supply']['overlap'].values() for v in overlaps)
 quoted=commit(quote,current);assert quoted['children'][0]['purchase_order_id']!=po['children'][0]['purchase_order_id']
 with e.connect() as c:assert c.execute(text("SELECT count(*) FROM purchasing_v4.po_ownership WHERE origin='V4'")).scalar_one()==2


def test_issued_request_survives_discard_of_its_source_worksheet(quote):
 e,state,rid,_=quote
 with e.begin() as c:
  current=worksheets.snapshot(c,state['session']['id'],1);worksheets.patch(c,actor_id=1,session_id=state['session']['id'],command_id=str(uuid4()),expected_version=current['session']['version'],changes=[],discard=True)
 r=final(quote);assert not r['blocking'];result=commit(quote,r);assert result['commitments_created']==1
 with e.connect() as c:
  assert worksheets.snapshot(c,state['session']['id'],1)['session']['state']=='DISCARDED'
  assert q.head(c,rid)['state']=='COMMITTED'
